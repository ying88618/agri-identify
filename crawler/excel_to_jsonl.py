# -*- coding: utf-8 -*-
"""
excel_to_jsonl.py — 「农业病虫害信息检索数据集」Excel → 结构化 JSONL

用法: python crawler/excel_to_jsonl.py
输入: E:\RAG_agri\ERAG_agriexcel\农业病虫害信息检索数据集\ (病害/虫害按作物分文件夹; 检疫/入侵按类型分)
输出: crawler/data/agri_pests.jsonl
    每行一条病虫害: {crop, category, name, source_file, fields:{字段:文本}}
"""
import json
import os

import openpyxl
import xlrd

SRC = r"E:\RAG_agri\ERAG_agriexcel\农业病虫害信息检索数据集"
OUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data")
OUT = os.path.join(OUT_DIR, "agri_pests.jsonl")
NAME_COL = "中文名称"


def read_rows(path):
    """读表头 + 所有数据行(以列名为键); 兼容 .xls / .xlsx"""
    rows = []
    if path.lower().endswith(".xls"):
        wb = xlrd.open_workbook(path)
        ws = wb.sheet_by_index(0)
        headers = [str(ws.cell_value(0, c)).strip() for c in range(ws.ncols)]
        for r in range(1, ws.nrows):
            rows.append({h: ws.cell_value(r, c) for c, h in enumerate(headers)})
    else:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        headers = [str(c.value).strip() if c.value else f"col{i}" for i, c in enumerate(ws[1])]
        for row in ws.iter_rows(min_row=2, values_only=True):
            rows.append({h: (v if v is not None else "") for h, v in zip(headers, row)})
        wb.close()
    return rows


def clean(v):
    if v is None:
        return ""
    return str(v).replace("\r", "").replace("\u3000", " ").strip()


def walk_crop_dir(root, cat, records):
    """农业病虫害: root/病害|虫害/<作物>/*.xls|xlsx"""
    cat_dir = os.path.join(root, "农业病虫害", cat)
    if not os.path.isdir(cat_dir):
        return
    for crop_dir in sorted(os.listdir(cat_dir)):
        crop_path = os.path.join(cat_dir, crop_dir)
        if not os.path.isdir(crop_path):
            continue
        for fn in os.listdir(crop_path):
            if not fn.lower().endswith((".xls", ".xlsx")):
                continue
            for row in read_rows(os.path.join(crop_path, fn)):
                name = clean(row.get(NAME_COL, ""))
                if not name:
                    continue
                fields = {h: clean(row.get(h, "")) for h in row if clean(row.get(h, ""))}
                fields.pop(NAME_COL, None)
                records.append({"crop": crop_dir, "category": cat, "name": name,
                                "source_file": f"{cat}/{crop_dir}/{fn}", "fields": fields})


def walk_type_dir(root, top, records):
    """检疫性物种/外来入侵物种: top/<类型>/*.xls|xlsx (无作物维度)"""
    top_dir = os.path.join(root, top)
    if not os.path.isdir(top_dir):
        return
    for sub in sorted(os.listdir(top_dir)):
        sub_path = os.path.join(top_dir, sub)
        if not os.path.isdir(sub_path):
            continue
        for fn in os.listdir(sub_path):
            if not fn.lower().endswith((".xls", ".xlsx")):
                continue
            for row in read_rows(os.path.join(sub_path, fn)):
                name = clean(row.get(NAME_COL, ""))
                if not name:
                    continue
                fields = {h: clean(row.get(h, "")) for h in row if clean(row.get(h, ""))}
                fields.pop(NAME_COL, None)
                records.append({"crop": None, "category": f"{top}-{sub}", "name": name,
                                "source_file": f"{top}/{sub}/{fn}", "fields": fields})


def main():
    records = []
    for cat in ("病害", "虫害"):
        walk_crop_dir(SRC, cat, records)
    for top in ("检疫性物种", "外来入侵物种"):
        walk_type_dir(SRC, top, records)

    os.makedirs(OUT_DIR, exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        for r in records:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")

    from collections import Counter
    print(f"总条目: {len(records)} 种病虫害")
    for cat, n in Counter(r["category"] for r in records).items():
        print(f"  {cat}: {n}")

    sample = next((x for x in records if x["crop"] == "水稻" and x["name"] == "稻瘟病"), None)
    if sample:
        print("\n样例(水稻/稻瘟病):")
        print(json.dumps(sample, ensure_ascii=False, indent=2)[:900])
    print(f"\n已导出 → {OUT}")


if __name__ == "__main__":
    main()
